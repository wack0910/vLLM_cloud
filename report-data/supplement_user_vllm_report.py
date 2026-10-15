from pathlib import Path
import ast, hashlib, json
from lxml import etree
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'보고서/vLLM_벤치마킹_결과_보고서_표중심_수정본.docx'
DEST=ROOT/'보고서/vLLM_벤치마킹_보고서_경어체_플래그설명추가.docx'
QA=ROOT/'보고서/vllm-supplement-qa';QA.mkdir(exist_ok=True)
source_hash=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
doc=Document(SOURCE)
original_paragraphs=list(doc.paragraphs)
original_xml=[etree.tostring(p._p) for p in original_paragraphs]
table_xml=[etree.tostring(t._tbl) for t in doc.tables]

# Only prose that exactly matches the assistant-generated version may change.
# All user-modified paragraphs, tables, drawings, headings and links are frozen.
tree=ast.parse((ROOT/'report-data/build_vllm_table_report.py').read_text(encoding='utf-8'))
assistant_prose={n.args[0].value for n in ast.walk(tree)
    if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id=='p'
    and n.args and isinstance(n.args[0],ast.Constant) and isinstance(n.args[0].value,str)}
replacements={
    '주요 변화  출력 처리량 132.74→132.78 token/s. 평균 TTFT와 P95 ITL도 거의 같아 성능 차이가 관측되지 않았다.':
    '주요 변화  출력 처리량은 132.74→132.78 token/s로 나타났습니다. 평균 TTFT와 P95 ITL도 거의 같아 성능 차이가 관측되지 않았습니다.',
    '주요 변화  Eager 강제 시 처리량 14.25% 감소, 평균 TPOT 16.56% 증가, P95 TTFT 15.25% 증가.':
    '주요 변화  Eager 실행을 강제했을 때 처리량은 14.25% 감소했으며, 평균 TPOT는 16.56%, P95 TTFT는 15.25% 증가했습니다.',
    '활용  평상시에는 false를 유지한다. 그래프 관련 오류나 호환성을 점검할 때 true를 검토한다. 실행 경로 전체의 비교이며 CUDA Graph 단독 효과는 분리하지 않았다.':
    '활용  평상시에는 false를 유지하시는 것이 좋습니다. CUDA Graph 관련 오류나 호환성을 점검하실 때 true를 검토하실 수 있습니다. 이번 결과는 실행 경로 전체를 비교한 값이며, CUDA Graph 단독 효과는 분리하지 않았습니다.',
    '주요 변화  처리량 46.11% 증가, 평균 TTFT 34.90% 감소. 같은 문서나 시스템 프롬프트를 반복하는 작업에 적합하다.':
    '주요 변화  처리량은 46.11% 증가했으며, 평균 TTFT는 34.90% 감소했습니다. 같은 문서나 시스템 프롬프트를 반복해서 사용하는 작업에 적합합니다.',
    '해석 범위  캐시가 생성 계산을 직접 생략하는 것은 아니다. TPOT 개선은 입력 처리와의 자원 경쟁 완화 가능성이 있다. 적중률은 미계측이며 사전 요청으로 캐시가 예열됐을 수 있다.':
    '해석 범위  캐시가 생성 계산을 직접 생략하지는 않습니다. TPOT는 입력 처리와의 자원 경쟁이 줄어 개선됐을 가능성이 있습니다. 적중률은 측정하지 않았으며, 사전 요청으로 캐시가 예열됐을 수 있습니다.',
    '적용 조건  Prefix ON과 토큰 예산 1024 조합은 미측정이다. 조건별 1회 결과이며 모델 품질·장시간 안정성의 검증은 아니다. Qwen3.8은 혼합 구조여서 일반 full-attention 모델의 KV 계산식을 그대로 적용하지 않는다. [5]':
    '적용 조건  Prefix ON과 토큰 예산 1024의 조합은 측정하지 않았습니다. 각 조건을 1회 측정한 결과이므로 모델 품질과 장시간 안정성까지 검증한 것은 아닙니다. Qwen3.8은 혼합 구조이므로 일반 full-attention 모델의 KV 계산식을 그대로 적용하실 수 없습니다. [5]',
    '처리량은 높을수록, 시간은 낮을수록 좋음 · 변화율은 반올림 전 원본값으로 계산':
    '처리량은 높을수록, 시간은 낮을수록 좋습니다. 변화율은 반올림 전 원본값으로 계산했습니다.',
}
changed=[]
for idx,p in enumerate(original_paragraphs):
    old=p.text
    if old in replacements:
        assert old in assistant_prose, ('User-edited paragraph cannot be changed',idx)
        new=replacements[old]
        # Keep the existing paragraph properties, including custom spacing and style.
        for child in list(p._p):
            if child.tag!=qn('w:pPr'):p._p.remove(child)
        if '  ' in new:
            lead,body=new.split('  ',1);p.add_run(lead+'  ').bold=True;p.add_run(body)
        else:p.add_run(new)
        changed.append(idx)

def font(run,size=10,bold=False):
    run.font.name='맑은 고딕';run.font.size=Pt(size);run.font.color.rgb=RGBColor(0,0,0);run.bold=bold
    run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
def paragraph(text,size=10):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.05
    font(p.add_run(text),size)
    return p
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for vals in rows:
        for c,txt in zip(t.add_row().cells,vals):c.text=txt
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bd.append(el)
            cp.append(bd)
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');mar.append(el)
            cp.append(mar)
            if ri==0:
                el=OxmlElement('w:shd');el.set(qn('w:fill'),'E6EDF2');cp.append(el)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.05
                for run in p.runs:font(run,9.2,ri==0)
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=.3

h=doc.add_paragraph(style='Heading 1');h.paragraph_format.page_break_before=True
font(h.add_run('9. 플래그 사용처와 환경별 선택'),16,True)
paragraph('서버 플래그는 vllm serve 또는 서버 YAML에 적용합니다. 변경 후 서버를 재시작하셔야 합니다. 아래 환경별 선택은 이번 측정에 따른 검토 방향이며, 모든 모델의 최적값을 의미하지는 않습니다.')
table(['플래그와 용도','유리한 환경과 선택 방향'],[
('max-num-seqs\n단계별 시퀀스 수 상한입니다.',
 '여러 사용자의 요청을 함께 처리하고 첫 응답 대기를 줄이실 때 큰 값을 검토하실 수 있습니다. 이번 결과에서는 32가 TTFT에, 16 또는 24가 TPOT에 유리했습니다.'),
('max-num-batched-tokens\n단계별 토큰 예산입니다.',
 '긴 입력과 스트리밍 생성이 함께 들어오는 채팅에서는 작은 예산이 생성 중 긴 멈춤을 줄이는 데 유리할 수 있습니다. 이번 환경에서는 1024의 P95 ITL이 낮았습니다.'),
('gpu-memory-utilization\n실행기 GPU 메모리 예산입니다.',
 'GPU를 단독 사용하시고 메모리 여유가 충분하면 큰 예산을 검토하실 수 있습니다. 다른 프로세스와 GPU를 공유하실 때는 낮은 예산이 메모리 여유 확보에 유리합니다.'),
('max-model-len\n입력과 출력 합계의 허용 상한입니다.',
 '긴 문서나 대화 이력을 처리하실 때 큰 값이 필요합니다. 현재 입력 길이에는 4096이면 충분합니다. 상한만 늘려도 짧은 요청이 빨라지는 것은 아닙니다.'),
('enforce-eager\nEager 실행을 강제하는 옵션입니다.',
 '일반 서빙에서는 false가 적합하며 이번 측정에서도 처리량에 유리했습니다. CUDA Graph 관련 오류를 분리하거나 호환성을 점검하실 때 true를 검토하실 수 있습니다.'),
('enable-prefix-caching\n공통 접두사의 입력 계산을 재사용합니다.',
 '같은 시스템 프롬프트, 문서 또는 대화 이력을 반복해서 보내는 환경에서 ON이 유리합니다. 공통 접두사가 없거나 생성 시간이 대부분인 작업에서는 효과가 제한적일 수 있습니다.'),
],[2.55,4.4])
paragraph('공통 벤치마크 플래그는 vllm bench serve에 적용합니다. 서버 설정 비교에서는 같은 값을 유지하셔야 합니다.')
table(['플래그','사용처와 적합한 환경'],[
('--max-concurrency / --request-rate',
 '클라이언트 부하를 조절합니다. 동시성 32와 inf는 요청을 가능한 빠르게 보내는 부하 시험에 적합합니다. 실제 도착률을 재현하시려면 유한한 request-rate로 별도 측정하셔야 합니다.'),
('--dataset-name / 입력·출력 길이',
 'random은 길이를 맞춘 설정 비교에, prefix_repetition은 반복 접두사 캐시 검증에 적합합니다. 실제 채팅 품질을 평가하는 데이터셋은 아닙니다.'),
('--num-prompts / --seed',
 '측정 규모와 입력 재현 조건을 정합니다. 100개와 seed=0을 유지하여 실행 부담을 줄이고 비교 조건을 맞췄습니다. 1회 측정으로 장기 안정성을 확정하실 수는 없습니다.'),
],[2.55,4.4])
paragraph('표 확인 안내  기존 표의 default는 공식 vLLM 기본값이 아니라 이번 비교의 기준값 16과 2048을 뜻합니다. 3절의 평균 토큰당 시간(sec) TPOT에 적힌 145.02·146.38·146.05는 원본 JSON 기준 ms 값입니다. 해당 표기는 수정하지 않고 이 안내로 구분했습니다.',9)

doc.save(DEST)
check=Document(DEST)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
assert len(check.paragraphs)>=len(original_paragraphs)
for idx,old in enumerate(original_xml):
    if idx not in changed:
        assert etree.tostring(check.paragraphs[idx]._p)==old, ('Frozen paragraph changed',idx)
for idx,old in enumerate(table_xml):
    assert etree.tostring(check.tables[idx]._tbl)==old, ('Existing table changed',idx)
assert len(doc.inline_shapes)==len(check.inline_shapes)
(QA/'preservation.json').write_text(json.dumps({
    'source_sha256':source_hash,'source_unchanged':True,
    'original_paragraph_count':len(original_paragraphs),
    'assistant_prose_paragraphs_updated':changed,
    'preserved_paragraph_count':len(original_paragraphs)-len(changed),
    'preserved_existing_table_count':len(table_xml),
    'preserved_inline_images':len(check.inline_shapes),
},ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved supplement. Frozen paragraphs and all original tables verified.')

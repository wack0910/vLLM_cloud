from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'양자화 자료';OUT=ROOT/'보고서'
QA=OUT/'quantization-practice-qa';QA.mkdir(exist_ok=True)
result=json.loads((SRC/'08_fp8_inference.json').read_text(encoding='utf-8-sig'))
assert result['choices'][0]['message']['content']=='4826'
recipe=(SRC/'02_fp8_recipe.py').read_text(encoding='utf-8-sig')
assert 'scheme="FP8_DYNAMIC"' in recipe
assert 'Qwen2.5-7B-FP8-oneshot' in recipe

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for name in ['Normal','Title','Heading 1','Caption']:
    st=doc.styles[name];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(21)
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
doc.styles['Caption'].font.size=Pt(9)
doc.core_properties.title='Qwen2.5 7B FP8 양자화 실습 정리'
doc.core_properties.author=''

def p(text):return doc.add_paragraph(text)
def h(text,new=False):
    hp=doc.add_heading(text,1)
    if new:hp.paragraph_format.page_break_before=True
def code(text):
    pp=doc.add_paragraph();pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.space_after=Pt(9)
    pp.paragraph_format.line_spacing=1.0
    rr=pp.add_run(text);rr.font.name='Consolas';rr.font.size=Pt(9)
    rr._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
def shot(file,width,caption):
    cp=doc.add_paragraph(caption,style='Caption');cp.alignment=1;cp.paragraph_format.keep_with_next=True
    pp=doc.add_paragraph();pp.alignment=1;pp.paragraph_format.space_after=Pt(7)
    pp.add_run().add_picture(str(SRC/file),width=Inches(width))
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=txt
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bd.append(el)
            cp.append(bd)
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');mar.append(el)
            cp.append(mar)
            if ri==0:
                el=OxmlElement('w:shd');el.set(qn('w:fill'),'E6EDF2');cp.append(el)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2)
                pp.paragraph_format.line_spacing=1.05
                for r in pp.runs:r.font.size=Pt(9.5);r.bold=ri==0
    sp=doc.add_paragraph();sp.paragraph_format.space_after=Pt(0);sp.paragraph_format.line_spacing=.3

doc.add_paragraph('Qwen2.5 7B FP8 양자화 실습 정리',style='Title')
p('Qwen2.5-7B-Instruct를 FP8로 양자화하고, 저장한 모델을 vLLM으로 실행한 뒤 API로 응답을 확인했습니다.')

h('1. 원본 모델 준비')
p('Qwen/Qwen2.5-7B-Instruct와 L40S 1장을 사용했습니다. 양자화에는 LLM Compressor를, 모델 실행에는 vLLM을 사용했습니다.')
p('로컬에 준비한 원본 모델과 토크나이저를 불러왔습니다. 원본 경로는 models/Qwen2.5-7B-Instruct이며, 모델은 BF16으로 GPU에 로드했습니다.')

h('2. FP8 양자화 실행')
shot('양자화 설정값.png',2.9,'이번 실습에 적용한 FP8 설정')
p('FP8_DYNAMIC으로 선형 계층을 양자화했습니다. lm_head와 임베딩은 제외 대상으로 지정했으며, 별도 보정 데이터나 추가 학습 없이 진행했습니다.')
p('oneshot에 모델·토크나이저·양자화 설정·저장 경로를 전달했습니다. 아래는 실행 스크립트의 핵심 호출 부분입니다.')
code('oneshot(model=model, tokenizer=tokenizer,\n        recipe=recipe, output_dir=str(OUTPUT_DIR),\n        log_dir=str(LOG_DIR))')
p('준비한 스크립트는 다음 명령으로 실행했습니다.')
code('cd /root/workspace/quantization\npython scripts/quantize_fp8_oneshot.py')

h('3. 양자화 모델 저장 확인')
table(['항목','확인 내용'],[
    ('저장 폴더','outputs/Qwen2.5-7B-FP8-oneshot'),
    ('완료 확인','FP8 quantization completed 메시지를 확인했습니다.'),
],[1.4,5.55])
p('이후 vLLM에서 원본 모델 경로 대신 저장한 FP8 모델 폴더를 지정하여 사용했습니다.')

h('4. vLLM으로 모델 실행',True)
p('서빙 설정 파일 config/fp8-serve.yaml에서 FP8 모델 경로와 API 모델 이름을 지정했습니다. 아래는 모델 연결에 필요한 핵심 설정입니다.')
code('model: /root/workspace/quantization/outputs/Qwen2.5-7B-FP8-oneshot\nserved-model-name: qwen-fp8\nhost: "127.0.0.1"\nport: 8000')
p('준비한 설정 파일로 서버를 실행했습니다. vLLM은 저장된 양자화 설정을 읽어 모델을 로드했습니다.')
code('vllm serve --config config/fp8-serve.yaml')
p('Application startup complete 메시지로 서버 기동을 확인했습니다. API 주소는 http://127.0.0.1:8000이며, 요청할 때는 qwen-fp8이라는 이름을 사용합니다.')

h('5. API로 모델 사용')
p('서버를 켜 둔 상태에서 같은 서버의 다른 세션에서 아래 요청을 보냈습니다.')
code('curl -sS --fail-with-body \\\n  http://127.0.0.1:8000/v1/chat/completions \\\n  -H "Content-Type: application/json" \\\n  -d \'{\n    "model": "qwen-fp8",\n    "messages": [{"role": "user",\n      "content": "127 곱하기 38의 결과만 숫자로 답해줘."}],\n    "temperature": 0,\n    "max_tokens": 80\n  }\'')
shot('추론 검증.png',6.7,'실제 API 요청과 모델 응답')
table(['요청','응답'],[
    ('127 곱하기 38의 결과만 숫자로 답해줘.','4826'),
],[5.4,1.55])
p('양자화 모델이 API 요청을 받아 답변하는 것을 확인했습니다. 이후에는 messages의 질문을 바꾸어 같은 방식으로 사용하실 수 있습니다.')
p('오늘은 원본 모델 준비 → FP8 양자화 → 모델 저장 → vLLM 실행 → API 호출까지 실습했습니다.')

doc.save(OUT/'Qwen2.5_7B_FP8_양자화_실습정리.docx')
print('Saved concise quantization practice summary')

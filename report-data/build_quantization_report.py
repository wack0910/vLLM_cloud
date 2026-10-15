from pathlib import Path
import json, re, sys, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'양자화 자료'
OUT=ROOT/'보고서';OUT.mkdir(exist_ok=True)
QA=OUT/'quantization-qa';QA.mkdir(exist_ok=True)
read=lambda name:(SRC/name).read_text(encoding='utf-8-sig')
config=json.loads(read('05_fp8_config.txt'))
response=json.loads(read('08_fp8_inference.json'))
models=json.loads(read('08_fp8_models.json'))
sizes=read('04_size_comparison.txt')
bf16=float(re.search(r'BF16 weights: ([\d.]+)',sizes)[1])
fp8=float(re.search(r'FP8 weights: ([\d.]+)',sizes)[1])
reported=float(re.search(r'Storage reduction: ([\d.]+)',sizes)[1])
assert abs((1-fp8/bf16)*100-reported)<.01
assert config['quant_method']=='compressed-tensors' and config['quantization_status']=='compressed'
group=config['config_groups']['group_0']
assert group['weights']['num_bits']==8 and group['input_activations']['num_bits']==8
assert response['choices'][0]['message']['content']=='4826' and 127*38==4826
assert response['usage']==dict(prompt_tokens=48,total_tokens=53,completion_tokens=5,prompt_tokens_details=None,completion_tokens_details=None)
assert models['data'][0]['id']=='qwen-fp8'
log=read('07_fp8_vllm_memory.txt')
assert 'enable_prefix_caching=True' in log
assert '8.18 GiB memory' in log and '565,520 tokens' in log

sys.path.insert(0,str(ROOT/'.plot-deps'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
plt.rcParams.update({'font.family':'Malgun Gothic','axes.unicode_minus':False})
fig,ax=plt.subplots(figsize=(8,2.3),layout='constrained')
bars=ax.barh(['BF16 원본','FP8 양자화'],[bf16,fp8],color=['#9a9fa6','#35648a'],height=.48)
ax.invert_yaxis();ax.bar_label(bars,fmt='%.3f GiB',padding=6)
ax.set_xlim(0,17.1);ax.set_xlabel('safetensors 가중치 파일 크기  GiB')
ax.set_title('가중치 저장 용량 42.82% 절감',fontsize=12)
ax.spines[['top','right']].set_visible(False)
fig.savefig(QA/'weight-size.png',dpi=200);plt.close(fig)

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.6)
sec.left_margin=sec.right_margin=Inches(.65)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.07
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(21)
doc.styles['Title'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
doc.styles['Caption'].font.size=Pt(9)
doc.core_properties.title='Qwen2.5 7B FP8 양자화 및 vLLM 서빙 결과 보고서'
doc.core_properties.author=''

def p(text):
    pp=doc.add_paragraph()
    if '  ' in text:
        a,b=text.split('  ',1);pp.add_run(a+'  ').bold=True;pp.add_run(b)
    else:pp.add_run(text)
    return pp
def h(text,new=False):
    pp=doc.add_heading(text,1)
    if new:pp.paragraph_format.page_break_before=True
def image(name,width=6.7,caption=None):
    if caption:
        cp=doc.add_paragraph(caption,style='Caption');cp.alignment=1;cp.paragraph_format.keep_with_next=True
    pp=doc.add_paragraph();pp.alignment=1;pp.paragraph_format.space_after=Pt(7)
    path=SRC/name if (SRC/name).exists() else QA/name
    pp.add_run().add_picture(str(path),width=Inches(width))
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=widths or [6.9/len(headers)]*len(headers)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for vals in rows:
        for c,txt in zip(t.add_row().cells,vals):c.text=str(txt)
    for ri,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if ri==0:tr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            pr.append(b)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'45');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            if ri==0:
                el=OxmlElement('w:shd');el.set(qn('w:fill'),'E6EDF2');pr.append(el)
            for pp in c.paragraphs:
                pp.alignment=0 if ci==0 else 1
                pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=1.05
                for run in pp.runs:run.font.size=Pt(9.5);run.bold=ri==0
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=.4

doc.add_paragraph('Qwen2.5 7B FP8 양자화 및\nvLLM 서빙 결과 보고서',style='Title')
p('실험일 2026년 10월 15일 · NVIDIA L40S 1장')
p('Qwen2.5-7B-Instruct의 BF16 가중치를 FP8로 변환하여 파일 크기를 42.82% 줄였다. 생성한 체크포인트를 vLLM으로 로드하고 Chat Completions 응답까지 확인했다. 이번 결과는 양자화·저장·서빙 경로의 동작 검증이며, 속도 향상과 정확도 유지율을 측정한 비교 실험은 아니다.')
h('1. 서버 환경과 실험 대상')
image('서버환경.png',6.2,'실행 환경 캡처')
p('환경  L40S 명목 48GB 1장, CPU 16 cores, RAM 명목 32GB 환경이다. nvidia-smi에서 확인한 GPU 메모리는 46,068 MiB이며, 초기 점유는 14 MiB였다. 모델은 기존 벤치마크의 Qwen3.8-27B-FP8과 별개인 Qwen/Qwen2.5-7B-Instruct를 사용했다.',)
p('소프트웨어  Python 3.10.12, PyTorch 2.14.0+cu130, Transformers 5.17.0, LLM Compressor 0.14.0, compressed-tensors 0.19.0, vLLM 0.31.0이다. 드라이버는 595.99.02이다. nvidia-smi의 CUDA 13.2 표시는 드라이버가 지원하는 CUDA 버전이며, PyTorch에서 확인한 CUDA는 13.0이다.')

h('2. FP8 양자화 설정과 실행 결과',True)
p('방식  LLM Compressor의 oneshot으로 학습이 끝난 모델에 양자화를 적용했다. BF16 모델을 cuda:0에 로드한 뒤 FP8_DYNAMIC을 적용하고 로컬 체크포인트로 저장했다. 별도 보정 데이터셋이나 추가 학습은 사용하지 않았다. [1]')
image('양자화 설정값.png',3.2,'QuantizationModifier 설정')
p('옵션 의미  targets="Linear"는 선형 계층을 대상으로 한다. FP8_DYNAMIC은 가중치를 채널별 정적 FP8로, 입력 활성화를 토큰별 동적 FP8로 양자화한다. lm_head와 임베딩은 제외 대상으로 지정했다. 임베딩은 Linear 대상이 아니며 저장된 ignore 목록에는 lm_head만 남아 있다. [1]')
image('양자화 결과.png',6.9,'양자화 및 체크포인트 저장 완료')
p('결과  FP8 quantization completed와 결과 저장 경로를 확인했다. 저장된 설정은 quant_method=compressed-tensors, quantization_status=compressed이며 가중치와 입력 활성화 모두 8비트이다. 기존 결과 덮어쓰기를 막는 검사도 스크립트에 포함했다.')

h('3. 양자화 과정의 시간과 자원 사용',True)
table(['항목','기록값'],[
    ['스크립트 내부 경과 시간','18.93초'],
    ['프로세스 전체 wall clock','23.19초'],
    ['PyTorch allocated 최고값','14,957.36 MiB'],
    ['PyTorch reserved 최고값','15,086.00 MiB'],
    ['샘플링한 GPU 전체 메모리 최고값','15,602 MiB'],
    ['CPU 메모리 최대 RSS','15,857,080 KiB ≈ 15.12 GiB'],
],[3.95,2.95])
image('nvidia-smi 결과.png',6.3,'14 MiB를 점유한 유휴 GPU 화면  서빙 후 점유량 아님')
p('해석  18.93초에는 모델 로드와 결과 저장도 포함된다. 23.19초는 프로세스 시작·종료까지 포함하는 외부 측정이다. 세 메모리 최고값은 측정 범위가 다르며 합산하지 않는다. GPU 샘플링 최고값은 관측 사이의 순간 최고값을 놓칠 수 있다.')
p('메모리 구분  allocated는 PyTorch 텐서에 실제 할당된 메모리, reserved는 PyTorch 할당기가 확보한 메모리다. nvidia-smi는 GPU 전체 점유를 관측한다. CPU RSS는 시스템 RAM 사용량이며 GPU VRAM과 별개이다.')
p('관측 결과  양자화 직전 13시 46분과 직후 13시 48분의 별도 기록은 모두 14 MiB이다. 프로세스 종료 후 GPU 메모리가 반환된 상태이며, 과정 중 관측 최고값 15,602 MiB와 모순되지 않는다. 낮은 종료 후 점유량을 양자화 실행 중 사용량으로 해석하지 않는다.')

h('4. 모델 크기 절감 효과',True)
image('모델 크기 비교.png',6.9,'원본과 FP8 체크포인트 파일')
image('크기 절감율.png',5.6,'가중치 파일 크기 합계와 절감률 계산')
image('weight-size.png',6.3,'가중치 파일 용량 비교')
table(['지표','BF16','FP8'],[
    ['safetensors 가중치 합계','14.185 GiB','8.111 GiB'],
    ['가중치 파일 수','4개','1개'],
    ['원본 대비 용량','100%','57.18%'],
],[3.3,1.8,1.8])
p('결과  가중치 파일은 약 6.074 GiB 줄었으며, 절감률은 (1−FP8/BF16)×100=42.82%이다. 폴더 크기 15G→8.2G는 다른 파일과 반올림이 포함된 값이므로, 위 비교는 safetensors 파일 합계를 기준으로 했다.')
p('해석과 활용  16비트에서 8비트로 바꾸어도 모든 계층이 FP8이 되는 것은 아니다. 양자화에서 제외된 가중치와 스케일 등으로 정확히 50% 줄지 않는 것이 설명된다. 저장·전송 용량을 줄이는 데 유리하며, L40S에서 BF16 원본도 로드 가능했으므로 이번 실험은 실행 불가 모델을 실행 가능하게 만든 사례는 아니다.')

h('5. 양자화 모델의 vLLM 서빙 설정',True)
image('Serving 값.png',5.3,'vLLM 실행용 YAML')
p('양자화와 서빙의 관계  oneshot은 양자화 방식과 체크포인트를 저장한다. max-model-len, max-num-seqs 등의 실행 옵션은 별도 YAML에서 지정한다. 이번 코드에서 서빙 플래그를 양자화 레시피에 내장한 것은 아니다. 저장된 quant_method=compressed-tensors와 compressed 상태를 vLLM 기동 로그에서도 확인했다.')
table(['옵션','설정값','역할'],[
    ['served-model-name','qwen-fp8','API 요청에서 사용할 이름'],
    ['tensor-parallel-size','1','GPU 1장으로 실행'],
    ['gpu-memory-utilization','0.9','실행기 GPU 메모리 예산 비율'],
    ['max-model-len','4096','입력과 출력 합계의 허용 상한'],
    ['max-num-seqs','24','단계별 시퀀스 처리 상한'],
    ['max-num-batched-tokens','1024','단계별 토큰 처리 예산'],
    ['enable-chunked-prefill','true','긴 입력의 분할 처리'],
    ['dtype / kv-cache-dtype','auto / auto','비양자화 연산 및 KV 정밀도 자동 선택'],
    ['enforce-eager','false','CUDA Graph 활용 경로 허용'],
],[2.8,1.2,2.9])
p('적용 차이  YAML의 enable-prefix-caching은 false지만 실제 엔진 설정은 enable_prefix_caching=True이다. 따라서 이번 서빙 결과를 Prefix OFF 조건으로 기록하지 않는다. OFF가 필요하면 설치 버전의 명시적 --no-enable-prefix-caching 사용 후 적용값을 확인해야 한다. 원인은 이 자료만으로 확정하지 않는다.')

h('6. 서빙 기동과 GPU 메모리 결과',True)
image('llm serving 결과.png',6.3,'API 서버 시작 및 application startup complete')
table(['지표','기록값'],[
    ['FP8 체크포인트 크기','8.11 GiB'],
    ['모델 로드 시 메모리','8.18 GiB'],
    ['모델 로드 경과 시간','4.25초'],
    ['초기화 시간  프로파일·KV 생성·워밍업','35.62초'],
    ['KV Cache에 사용 가능한 메모리','30.2 GiB'],
    ['GPU KV Cache 토큰 용량','565,520토큰'],
    ['4096토큰 기준 캐시 수용량 추정','138.07배'],
    ['서빙 후 GPU 전체 메모리 점유','40,308 MiB ≈ 39.36 GiB'],
    ['VLLM EngineCore 프로세스 메모리','40,286 MiB'],
],[4.25,2.65])
p('해석  가중치가 8.11 GiB인데 GPU가 약 39.36 GiB 점유되는 것은 캐시와 실행용 메모리도 확보하기 때문이다. gpu-memory-utilization=0.9 조건에서 절감한 가중치 용량을 캐시 공간으로 활용할 수 있으므로, 파일 절감률과 GPU 전체 점유 절감률을 동일하게 보지 않는다. 30.2 GiB는 캐시 공간이며 실제 요청이 차지한 KV 사용률은 아니다.')
p('용량과 동시성  138.07배는 565,520÷4096으로 산출한 캐시 수용량 추정이다. 실제 동시 처리 성능이나 처리량을 의미하지 않으며, max-num-seqs=24라는 별도 처리 상한도 있다. FP8 가중치와 KV Cache 양자화는 별개이고, 저장된 kv_cache_scheme은 null이다.')
p('비교 범위  위 서빙 점유량은 서빙 후 텍스트 기록에서 확인했다. BF16을 같은 조건으로 서빙한 기록이 없어 GPU 전체 메모리 절감률은 계산하지 않는다.')

h('7. 추론 검증과 종합 결과',True)
image('추론 검증.png',6.9,'Chat Completions 요청 및 응답')
table(['검증 항목','결과'],[
    ['API 모델 이름','qwen-fp8'],
    ['등록된 모델 경로','outputs/Qwen2.5-7B-FP8-oneshot'],
    ['요청','127 곱하기 38의 결과만 숫자로 응답'],
    ['반환 내용','4826  정답과 일치'],
    ['입력 / 출력 / 합계 토큰','48 / 5 / 53'],
    ['종료 사유','stop'],
],[2.6,4.3])
p('검증 결과  양자화 체크포인트를 vLLM이 로드하고 API로 올바른 계산 응답을 반환했다. FP8 여부는 모델의 자기소개가 아니라 quantization_config의 8비트 설정, compressed 상태, 실제 로드 로그로 확인했다. API 별칭 qwen-fp8 자체만으로 양자화를 증명하지는 않는다.')
p('정상 여부와 한계  체크포인트 저장 완료, compressed-tensors 로드, 서버 시작 완료, 정답 응답이라는 기본 동작 검증을 통과했다. 42.82% 파일 절감과 높은 서빙 메모리 점유는 선택한 계층 및 캐시 예산으로 설명된다. 다만 산술 문제 1개는 전반적인 정확도 유지 증거가 아니며, TTFT·TPOT·Throughput과 BF16 대비 속도 향상도 미측정이다.')
p('적합한 활용  이 FP8_DYNAMIC 경로는 별도 보정 데이터를 준비하지 않고 FP8 체크포인트를 만들어 vLLM으로 배포하려는 경우에 활용할 수 있다. 실제 운영 도입 판단에는 같은 모델·서빙 조건에서 BF16과 FP8의 부하 시험 및 과제별 정확도 비교가 추가로 필요하다. AWQ·GPTQ·Unsloth와의 실측 비교는 이번 실험 범위에 포함되지 않는다.')
p('결론  가중치 저장 용량 42.82% 절감과 단일 GPU에서의 서빙·추론 동작을 확인했다. 양자화 단계와 서빙 설정을 분리해 재현 가능한 실행 경로를 확보했으며, Prefix Caching의 실제 적용값은 YAML과 달리 ON으로 기록한다.')

refs=[
    ('[1] LLM Compressor FP8 공식 예제','https://docs.vllm.ai/projects/llm-compressor/en/latest/examples/quantization_w8a8_fp8/'),
    ('[2] vLLM 서버 옵션','https://docs.vllm.ai/en/stable/cli/serve/'),
    ('[3] Qwen2.5 모델 카드','https://huggingface.co/Qwen/Qwen2.5-7B-Instruct'),
    ('준비 과정 참고  ThakiCloud 양자화 글','https://thakicloud.com/tech-blog/ko/llmops/llm-quantization-vllm-serving-unsloth/'),
]
pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(3)
for idx,(label,url) in enumerate(refs):
    if idx:pp.add_run('  ·  ')
    rel=pp.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'17');pr.append(sz);rr.append(pr)
    tx=OxmlElement('w:t');tx.text=label;rr.append(tx);link.append(rr);pp._p.append(link)

path=OUT/'Qwen2.5_7B_FP8_양자화_서빙_결과_보고서.docx'
doc.save(path)
(QA/'verified-summary.json').write_text(json.dumps({
    'bf16_weights_GiB':bf16,'fp8_weights_GiB':fp8,'reported_reduction_percent':reported,
    'recomputed_reduction_from_rounded_sizes_percent':(1-fp8/bf16)*100,
    'prefix_config_requested':False,'prefix_engine_effective':True,
    'quantization':config,'inference_usage':response['usage'],
    'source_sha256':{f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in SRC.iterdir() if f.is_file()},
},ensure_ascii=False,indent=2),encoding='utf-8')
print(str(path))

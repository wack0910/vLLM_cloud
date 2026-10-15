from pathlib import Path
import json,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'벤치마킹 결과';OUT=ROOT/'보고서'
QA=OUT/'vllm-table-qa';QA.mkdir(exist_ok=True)
D={p.name:json.loads((p/'benchmark.json').read_text(encoding='utf-8-sig')) for p in SRC.iterdir() if p.is_dir()}
B='base-r3072-o256-c32';S='seq24-r3072-o256-c32';T='seq32-r3072-o256-c32'
BT1='seq24-bt1024-r3072-o256-c32';BT4='seq24-bt4096-r3072-o256-c32'
G='seq24-gpu08-r3072-o256-c32';L='seq24-len8192-r3072-o256-c32';E='seq24-eager-r3072-o256-c32'
PO='prefix-off-p2048-s1024-o256-c32';PN='prefix-on-p2048-s1024-o256-c32'
for d in D.values():
    assert d['completed']==100 and d['failed']==0 and d['total_output_tokens']==25600
    assert d['max_concurrency']==32

sys.path.insert(0,str(ROOT/'.plot-deps'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
plt.rcParams.update({'font.family':'Malgun Gothic','axes.unicode_minus':False,'font.size':9})
chart_ids=[B,S,T,BT1,BT4,G,L,E]
chart_labels=['요청 수 16','요청 수 24','요청 수 32','토큰 예산 1024','토큰 예산 4096','메모리 예산 0.8','문맥 길이 8192','Eager 실행']
fig,axs=plt.subplots(1,2,figsize=(11,3.6),layout='constrained')
for ax,key,title,div in [(axs[0],'output_throughput','출력 처리량  token/s  높을수록 좋음',1),(axs[1],'p95_ttft_ms','P95 첫 응답 시간  초  낮을수록 좋음',1000)]:
    vals=[D[id][key]/div for id in chart_ids]
    bars=ax.barh(chart_labels,vals,color=['#939ba5','#35648a','#6195b2','#35648a','#939ba5','#939ba5','#939ba5','#939ba5'])
    ax.invert_yaxis();ax.set_title(title,fontsize=10);ax.spines[['top','right']].set_visible(False)
    ax.set_xlim(0,max(vals)*1.23);ax.bar_label(bars,fmt='%.1f',padding=3,fontsize=9)
fig.savefig(QA/'comparison.png',dpi=180);plt.close(fig)

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.07
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(5)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].paragraph_format.space_after=Pt(3)
doc.core_properties.title='vLLM 설정별 성능 비교 보고서'
doc.core_properties.author=''

def p(txt):
    pp=doc.add_paragraph()
    if '  ' in txt:
        lead,rest=txt.split('  ',1);pp.add_run(lead+'  ').bold=True;pp.add_run(rest)
    else:pp.add_run(txt)
    return pp
def h(txt,new=False):
    hp=doc.add_heading(txt,1)
    if new:hp.paragraph_format.page_break_before=True
def caption(txt):
    pp=doc.add_paragraph(txt,style='Caption');pp.paragraph_format.keep_with_next=True;return pp
def shots(items,w=3.05):
    caption('    |    '.join(label for _,label in items)).alignment=1
    pp=doc.add_paragraph();pp.alignment=1;pp.paragraph_format.space_after=Pt(7)
    for ix,(id,label) in enumerate(items):
        if ix:pp.add_run('   ')
        pp.add_run().add_picture(str(SRC/id/(id+'.png')),width=Inches(w))
def table(headers,rows,widths=None,size=9.4):
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    widths=widths or [6.95/len(headers)]*len(headers)
    for col,w in zip(tb.columns,widths):col.width=Inches(w)
    for c,txt in zip(tb.rows[0].cells,headers):c.text=txt
    for vals in rows:
        for c,txt in zip(tb.add_row().cells,vals):c.text=str(txt)
    for ri,row in enumerate(tb.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            cp.append(borders)
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');mar.append(el)
            cp.append(mar)
            if ri==0:
                el=OxmlElement('w:shd');el.set(qn('w:fill'),'E6EDF2');cp.append(el)
            for pp in c.paragraphs:
                pp.alignment=0 if ci==0 else 1
                pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.line_spacing=1.02
                for rr in pp.runs:rr.font.size=Pt(size);rr.bold=ri==0
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=.3

METRICS=[
('실행 시간  초','duration',1),
('출력 처리량  token/s','output_throughput',1),
('평균 첫 응답  초  TTFT','mean_ttft_ms',1000),
('P95 첫 응답  초  TTFT','p95_ttft_ms',1000),
('평균 토큰당 시간  ms  TPOT','mean_tpot_ms',1),
('P95 토큰당 시간  ms  TPOT','p95_tpot_ms',1),
('P95 스트림 간격  ms  ITL','p95_itl_ms',1),
('평균 완료 시간  초  E2EL','mean_e2el_ms',1000),
('실패 요청  개','failed',1),
]
def result_table(items,compare=None):
    headers=['비교 항목']+[label for _,label in items]
    if compare:headers.append(compare[2]+'\n변화율')
    rows=[]
    for label,key,div in METRICS:
        vals=[label]+[str(D[id][key]) if key=='failed' else f'{D[id][key]/div:,.2f}' for id,_ in items]
        if compare:
            a,b,_=compare
            vals.append('동일' if key=='failed' else f'{(D[b][key]/D[a][key]-1)*100:+.2f}%')
        rows.append(vals)
    if len(headers)==5:widths=[2.35,1.05,1.05,1.05,1.45]
    elif len(headers)==4:widths=[2.65,1.35,1.35,1.6]
    else:widths=[2.95,2,2]
    table(headers,rows,widths,9.1)
    small=p('처리량은 높을수록, 시간은 낮을수록 좋음 · 변화율은 반올림 전 원본값으로 계산')
    for r in small.runs:r.font.size=Pt(8.5)

doc.add_paragraph('vLLM 설정별 성능 비교 보고서',style='Title')
p('Qwen/Qwen3.8-27B-FP8 · NVIDIA L40S 48GB 1장 · 2026년 10월 14~15일 측정')
h('1. 테스트 환경과 기준 성능')
table(['항목','테스트 조건'],[
('장비','CPU 16 cores · RAM 32GB · L40S 1장'),
('측정 방식','공식 vllm bench serve · Chat Completions'),
('일반 부하','random · 입력 목표 3072토큰 · 출력 256토큰'),
('요청 조건','100개 · 동시성 32 · request-rate=inf · seed=0'),
('생성 조건','temperature=0 · ignore-eos · thinking OFF · warmups=0'),
('기준 설정','요청 수 상한 16 · 토큰 예산 2048 · GPU 메모리 예산 0.9'),
('기타 고정값','문맥 4096 · Prefix OFF · Chunked Prefill ON · Eager false · TP 1 · KV auto'),
],[1.4,5.55])
shots([(B,'기준 설정의 원본 결과')],2.85)
table(['실행 시간','출력 처리량','평균 TTFT','평균 TPOT','성공 / 실패'],[
('206.62초','123.90 token/s','32.06초','109.06ms','100 / 0')
],[1.2,1.65,1.35,1.35,1.4])
p('비교 기준  요청 수 실험은 16을, 그 외 일반 부하 실험은 24를 기준으로 비교한다. Prefix 실험은 별도 반복 입력 부하의 OFF와 ON을 비교한다.')
p('측정 범위  조건별 1회씩 총 10개 조건, 1,000개 요청 성공. 실제 일반 입력은 평균 3084.38토큰이며 출력은 모두 256토큰이다. warmups=0과 별도로 사전 단일 요청 테스트가 있다.')

h('2. 함께 처리할 요청 수를 늘리면',True)
p('조정 옵션  max-num-seqs = 16 / 24 / 32. 한 처리 단계의 시퀀스 상한이며, 클라이언트 동시성은 32로 유지했다. [1]')
shots([(S,'요청 수 상한 24'),(T,'요청 수 상한 32')])
result_table([(B,'16 기준'),(S,'24'),(T,'32')],(B,T,'16→32'))
p('주요 변화  32는 첫 응답이 빨라졌지만 평균 TPOT는 68.42% 증가했다. 24→32의 처리량 증가는 1.33%에 그쳤다.')
p('활용  첫 응답이 우선이면 32, 생성 중 응답성이 우선이면 16 또는 24를 검토한다. 24는 이번 조건의 절충 후보다.')

h('3. 한 번에 처리할 토큰 수를 바꾸면',True)
p('조정 옵션  max-num-batched-tokens = 1024 / 2048 / 4096. 입력 처리와 생성에 쓸 단계별 토큰 예산이다. Chunked Prefill은 ON이다. [1]')
shots([(BT1,'토큰 예산 1024'),(BT4,'토큰 예산 4096')])
result_table([(BT1,'1024'),(S,'2048 기준'),(BT4,'4096')],(S,BT1,'2048→1024'))
p('주요 변화  1024에서 P95 ITL은 827.70→428.62ms로 48.22% 감소했다. 4096에서는 1263.06ms로 증가했다.')
p('활용  긴 입력이 섞인 채팅에서 생성 도중 긴 멈춤을 줄이려면 1024가 후보다. 약 1~2%인 처리량 차이의 우위는 반복 측정이 필요하다.')

h('4. GPU 메모리 예산을 줄이면',True)
p('조정 옵션  gpu-memory-utilization = 0.9 / 0.8. vLLM 실행기가 사용할 GPU 메모리 예산 비율이며, KV Cache 사용률과는 다르다. [1, 2]')
shots([(G,'GPU 메모리 예산 0.8')],3.05)
result_table([(S,'0.9 기준'),(G,'0.8')],(S,G,'0.9→0.8'))
p('주요 변화  처리량 12.26% 감소, 평균 TTFT 50.27% 증가. 평균 TPOT는 낮아졌지만 P95 TPOT는 높아졌다.')
p('활용  GPU 단독 사용 시 이번 조건은 0.9가 유리했다. 다른 프로세스에 메모리를 남길 때 0.8을 검토한다. KV 사용률·preemption 미계측으로 캐시 부족 여부는 확정할 수 없다.')

h('5. 받을 수 있는 최대 문맥을 늘리면',True)
p('조정 옵션  max-model-len = 4096 / 8192. 입력과 출력 합계의 허용 상한이다. 실제 입력 약 3084토큰과 출력 256토큰은 고정했다. [2]')
shots([(L,'최대 문맥 길이 8192토큰')],3.05)
result_table([(S,'4096 기준'),(L,'8192')],(S,L,'4096→8192'))
p('주요 변화  출력 처리량 132.74→132.78 token/s. 평균 TTFT와 P95 ITL도 거의 같아 성능 차이가 관측되지 않았다.')
p('활용  현재 부하는 4096이면 충분하다. 입력과 출력 합계가 4096을 넘는 작업에 8192를 사용한다. 실제 8192토큰 요청의 성능이나 KV 소비량 증가를 검증한 결과는 아니다.')

h('6. 일반 실행과 Eager 실행을 비교하면',True)
p('조정 옵션  enforce-eager = false / true. false는 CUDA Graph 활용 경로를 허용하고 true는 Eager 실행을 강제한다. [2]')
shots([(E,'Eager 실행 강제  true')],3.05)
result_table([(S,'일반 실행\nfalse 기준'),(E,'Eager 실행\ntrue')],(S,E,'false→true'))
p('주요 변화  Eager 강제 시 처리량 14.25% 감소, 평균 TPOT 16.56% 증가, P95 TTFT 15.25% 증가.')
p('활용  평상시에는 false를 유지한다. 그래프 관련 오류나 호환성을 점검할 때 true를 검토한다. 실행 경로 전체의 비교이며 CUDA Graph 단독 효과는 분리하지 않았다.')

h('7. 반복되는 입력을 캐시로 재사용하면',True)
p('조정 옵션  Prefix Caching OFF / ON. 동일한 접두사의 입력 계산을 재사용한다. 공통 2048토큰 1종류 + 접미사 1024토큰, 출력 256토큰의 prefix_repetition 부하를 사용했다. [3]')
shots([(PO,'입력 재사용 OFF'),(PN,'입력 재사용 ON')])
result_table([(PO,'OFF 기준'),(PN,'ON')],(PO,PN,'OFF→ON'))
p('주요 변화  처리량 46.11% 증가, 평균 TTFT 34.90% 감소. 같은 문서나 시스템 프롬프트를 반복하는 작업에 적합하다.')
p('해석 범위  캐시가 생성 계산을 직접 생략하는 것은 아니다. TPOT 개선은 입력 처리와의 자원 경쟁 완화 가능성이 있다. 적중률은 미계측이며 사전 요청으로 캐시가 예열됐을 수 있다.')

h('8. 어떤 설정을 선택하면 좋을까',True)
doc.add_picture(str(QA/'comparison.png'),width=Inches(6.9))
caption('일반 random 부하 8개 조건 비교  요청 수 외 변경은 상한 24 기준  Prefix 실험 제외')
table(['우선 목표','선택 후보','측정 근거'],[
('첫 응답 대기 줄이기','요청 수 상한 32','평균 TTFT 11.51초  16 대비 64.09% 감소'),
('생성 중 긴 멈춤 줄이기','요청 수 24 · 토큰 예산 1024','P95 ITL 428.62ms  2048 대비 48.22% 감소'),
('GPU 단독 사용','메모리 예산 0.9','0.8에서 출력 처리량 12.26% 감소'),
('현재 입력 길이 처리','최대 문맥 4096','8192의 출력 처리량 차이 +0.03%'),
('일반 서빙','Eager false','true에서 출력 처리량 14.25% 감소'),
('반복 프롬프트 처리','Prefix Caching ON','출력 처리량 46.11% 증가'),
],[1.65,2.3,3.0],9.0)
caption('수치 해석과 과제 검증 범위')
table(['항목','확인된 값 또는 상태','판단'],[
('요청 실패','전체 1000개 중 0개','단기 동작 확인  장시간 안정성은 미검증'),
('일반 부하의 응답 지연','평균 TTFT 11.51~32.06초\n평균 E2EL 57.06~66.19초','즉각적인 채팅용으로는 느림\n긴 입력 · 동시성 32 조건의 결과'),
('외부 공개 생성 속도 [4]','37.1~37.7 token/s\n입력 4096 · 다른 엔진의 생성 구간','측정 구간과 엔진이 달라 직접 우열 비교 불가'),
('KV 용량과 기능별 효과','KV 사용률·블록 수·preemption 미계측\nPagedAttention·Batching 개별 효과 미분리','지연 원인을 KV 부족으로 확정할 수 없음'),
('TP와 양자화 비교','TP 1 · FP8만 측정','다중 GPU 및 AWQ/GPTQ 대비 우위 미검증'),
],[1.65,2.65,2.65],8.7)
p('적용 조건  Prefix ON과 토큰 예산 1024 조합은 미측정이다. 조건별 1회 결과이며 모델 품질·장시간 안정성의 검증은 아니다. Qwen3.8은 혼합 구조여서 일반 full-attention 모델의 KV 계산식을 그대로 적용하지 않는다. [5]')
refs=[('[1] vLLM 최적화','https://docs.vllm.ai/en/stable/configuration/optimization/'),('[2] 서버 옵션','https://docs.vllm.ai/en/stable/cli/serve/'),('[3] Prefix Caching','https://docs.vllm.ai/en/stable/features/automatic_prefix_caching/'),('[4] 공개 결과','https://github.com/mozilla-ai/exp-llama-benchy/blob/main/data/20260912-l40s/report/qwen3.8-27b-all-servers.md'),('[5] 모델 구조','https://huggingface.co/Qwen/Qwen3.8-27B-FP8')]
pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(0)
for ix,(label,url) in enumerate(refs):
    if ix:pp.add_run('  ·  ')
    rel=pp.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'16');pr.append(sz);rr.append(pr)
    tx=OxmlElement('w:t');tx.text=label;rr.append(tx);link.append(rr);pp._p.append(link)

# Keep every numerical headline tied to the original unrounded JSON values.
headlines={
    'seq32_tpot_increase':(T,B,'mean_tpot_ms',68.42),
    'seq24_to_32_throughput':(T,S,'output_throughput',1.33),
    'bt1024_itl_change':(BT1,S,'p95_itl_ms',-48.22),
    'gpu08_throughput_change':(G,S,'output_throughput',-12.26),
    'gpu08_ttft_change':(G,S,'mean_ttft_ms',50.27),
    'eager_throughput_change':(E,S,'output_throughput',-14.25),
    'eager_tpot_change':(E,S,'mean_tpot_ms',16.56),
    'eager_p95ttft_change':(E,S,'p95_ttft_ms',15.25),
    'prefix_throughput_change':(PN,PO,'output_throughput',46.11),
    'prefix_ttft_change':(PN,PO,'mean_ttft_ms',-34.90),
    'seq32_ttft_change':(T,B,'mean_ttft_ms',-64.09),
}
errors=[]
for label,(i,b,k,expected) in headlines.items():
    actual=(D[i][k]/D[b][k]-1)*100
    if abs(actual-expected)>.011:errors.append((label,expected,round(actual,2)))
if errors:raise ValueError(errors)
doc.save(OUT/'vLLM_벤치마킹_결과_보고서_표중심_수정본.docx')
print('Saved table-focused revision')

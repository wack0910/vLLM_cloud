from pathlib import Path
import json, sys, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'벤치마킹 결과'
OUT=ROOT/'보고서'; OUT.mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT/'.plot-deps'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
plt.rcParams.update({'font.family':'Malgun Gothic','axes.unicode_minus':False,'font.size':10})
data={p.name:json.loads((p/'benchmark.json').read_text(encoding='utf-8-sig')) for p in SRC.iterdir() if p.is_dir()}
ids=['base-r3072-o256-c32','seq24-r3072-o256-c32','seq32-r3072-o256-c32','seq24-bt1024-r3072-o256-c32','seq24-bt4096-r3072-o256-c32','seq24-gpu08-r3072-o256-c32','seq24-len8192-r3072-o256-c32','seq24-eager-r3072-o256-c32','prefix-off-p2048-s1024-o256-c32','prefix-on-p2048-s1024-o256-c32']
b,s,t,bt1,bt4,g,l,e,po,pn=ids
for d in data.values():
    assert d['completed']==100 and d['failed']==0 and d['max_concurrency']==32
    assert d['total_output_tokens']==25600 and sum(d['output_lens'])==25600
    assert sum(d['input_lens'])==d['total_input_tokens']
    assert abs(d['output_throughput']-25600/d['duration'])<1e-8

def change(i,base,key): return (data[i][key]/data[base][key]-1)*100
def delta(i,base,key):
    n=change(i,base,key)
    return f'{abs(n):.1f}% '+('증가' if n>=0 else '감소')
labels=['seq16','seq24','seq32','seq24\nbt1024','seq24\nbt4096','seq24\nGPU 0.8','seq24\nlen8192','seq24\nEager']
fig,axs=plt.subplots(1,2,figsize=(11,3.6),layout='constrained')
for ax,key,title,div in [(axs[0],'output_throughput','출력 처리량  token/s  높을수록 좋음',1),(axs[1],'p95_ttft_ms','P95 TTFT  초  낮을수록 좋음',1000)]:
    vals=[data[x][key]/div for x in ids[:8]]
    bars=ax.barh(labels,vals,color=['#939ba5','#35648a','#6195b2','#35648a','#939ba5','#939ba5','#939ba5','#939ba5'])
    ax.invert_yaxis(); ax.set_title(title,fontsize=11); ax.spines[['top','right']].set_visible(False)
    ax.set_xlim(0,max(vals)*1.23); ax.bar_label(bars,fmt='%.1f',padding=3,fontsize=9)
fig.savefig(OUT/'comparison.png',dpi=180); plt.close(fig)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.6); sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    style=doc.styles[name]; style.font.name='맑은 고딕'; style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(21)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Caption'].font.size=Pt(9)
doc.core_properties.title='vLLM 파라미터 벤치마킹 결과 보고서'
doc.core_properties.author=''

def p(text,boldlead=False):
    p=doc.add_paragraph()
    if boldlead and '  ' in text:
        lead,rest=text.split('  ',1); p.add_run(lead+'  ').bold=True;p.add_run(rest)
    else:p.add_run(text)
    return p
def heading(text):doc.add_heading(text,1)
def page(text):
    hp=doc.add_heading(text,1);hp.paragraph_format.page_break_before=True
def shots(items,width=3.15):
    cap=doc.add_paragraph('     |     '.join(label for _,label in items),style='Caption');cap.alignment=1;cap.paragraph_format.keep_with_next=True
    pp=doc.add_paragraph();pp.alignment=1
    for j,(id,label) in enumerate(items):
        if j:pp.add_run('   ')
        pp.add_run().add_picture(str(SRC/id/(id+'.png')),width=Inches(width))
    pp.paragraph_format.space_after=Pt(8)
def table(items):
    tb=doc.add_table(rows=1, cols=5);tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=['설정','출력 처리량\ntoken/s','평균 TTFT\n초','P95 TTFT\n초','평균 TPOT\nms']
    for c,txt in zip(tb.rows[0].cells,hdr):c.text=txt
    for id,label in items:
        d=data[id];vals=[label,f"{d['output_throughput']:.2f}",f"{d['mean_ttft_ms']/1000:.2f}",f"{d['p95_ttft_ms']/1000:.2f}",f"{d['mean_tpot_ms']:.2f}"]
        for c,txt in zip(tb.add_row().cells,vals):c.text=txt
    for ri,row in enumerate(tb.rows):
        trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(1.4);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'D9D9D9');borders.append(v)
            pr.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EDF2');pr.append(shade)
            for pp in c.paragraphs:
                pp.alignment=1;pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.space_after=Pt(4)
                for r in pp.runs:r.font.size=Pt(9);r.bold=ri==0
    p('') .paragraph_format.space_after=Pt(0)

doc.add_paragraph('vLLM 파라미터 벤치마킹 결과 보고서',style='Title')
p('Qwen/Qwen3.8-27B-FP8 · L40S 48GB 1장 · 측정일 2026년 10월 14~15일')
p('반복 접두사에서는 Prefix Caching이 출력 처리량을 46.1% 높였다. 시퀀스 수 증가는 첫 응답 대기를 줄였지만 토큰 생성 간격을 늘렸다. 따라서 최대 처리량만으로 설정을 선택하기보다 첫 응답과 스트리밍 지연을 함께 판단해야 한다.')
heading('1. 설정 값들과 기준 결과')
p('CPU 16 cores, RAM 32GB 환경에서 공식 vllm bench serve와 Chat Completions를 사용했다. random 입력 목표 3,072토큰, 출력 256토큰, 요청 100개, 클라이언트 동시성 32, request-rate=inf, seed=0, temperature=0, ignore-eos, thinking OFF를 적용했다.')
p('기준 설정은 max-num-seqs=16, max-num-batched-tokens=2048, gpu-memory-utilization=0.9, max-model-len=4096, Prefix OFF, enforce-eager=false이다. Chunked Prefill은 ON, TP=1, KV dtype=auto이다. 공식 기본값 전체를 뜻하지 않는다.')
shots([(b,'기준 설정 seq16')],3.0)
p('결과  100개 모두 성공했고 206.62초가 걸렸다. 출력 처리량은 123.90 token/s, 평균 TTFT는 32.06초, 평균 TPOT는 109.06ms였다. 입력은 템플릿을 포함해 평균 3,084.38토큰, 출력은 모두 256토큰이다.',True)
p('이후 단일 옵션 실험은 seq24를 비교 기준으로 삼는다. Prefix 실험만 별도 반복 접두사 부하로 OFF와 ON을 비교한다. 각 조건은 1회 측정이며 총 10개 조건에서 1,000개 요청이 성공했다. 서버 설정은 제출한 조건명과 설정 파일에 따른다.',False)

page('2. max-num-seqs 16 24 32 비교')
p('기능  한 스케줄링 단계에서 처리할 시퀀스 수의 상한이다. 값을 높이면 더 많은 요청이 함께 진행될 수 있지만, 각 요청에 배분되는 연산 자원과 캐시·상태 메모리 부담도 달라진다. 클라이언트 동시성 32와는 별도 설정이다. [1]',True)
shots([(s,'max-num-seqs=24'),(t,'max-num-seqs=32')])
table([(b,'16 기준'),(s,'24'),(t,'32')])
p(f'결과  16→24에서 처리량은 {delta(s,b,"output_throughput")}, 평균 TTFT는 {delta(s,b,"mean_ttft_ms")}했지만 평균 TPOT는 {delta(s,b,"mean_tpot_ms")}했다. 16→32에서는 평균 TTFT가 {delta(t,b,"mean_ttft_ms")}하고 평균 TPOT는 {delta(t,b,"mean_tpot_ms")}했다.',True)
p(f'해석과 활용  24→32의 처리량 증가는 {change(t,s,"output_throughput"):.1f}%에 그쳤고 P95 TPOT는 160.87→247.37ms로 악화됐다. 첫 응답이 중요하면 32, 생성 중 응답성을 중시하면 16 또는 24를 검토한다. 24는 이번 실험의 절충 후보이며 모든 지표의 최적값은 아니다.',True)

page('3. max-num-batched-tokens 1024 2048 4096 비교')
p('기능  한 스케줄링 단계의 토큰 처리 예산이다. Chunked Prefill을 켜면 긴 입력을 나누어 처리한다. 작은 예산은 생성 도중 긴 입력 계산이 끼어드는 시간을 줄일 수 있고, 큰 예산은 입력을 더 많이 처리할 수 있다. 실제 효과는 부하에 따라 달라진다. [1]',True)
shots([(bt1,'토큰 예산 1024'),(bt4,'토큰 예산 4096')])
table([(bt1,'1024'),(s,'2048 기준'),(bt4,'4096')])
p(f'결과  2048→1024에서 처리량은 {delta(bt1,s,"output_throughput")}, 평균 TTFT는 {delta(bt1,s,"mean_ttft_ms")}했다. 가장 큰 변화는 P95 ITL 827.70→428.62ms로 {abs(change(bt1,s,"p95_itl_ms")):.1f}% 감소한 점이다. 4096에서는 P95 ITL이 1,263.06ms로 {change(bt4,s,"p95_itl_ms"):.1f}% 증가했다.',True)
p('해석과 활용  이번 긴 입력·동시 요청 부하에서는 1024가 생성 중 긴 멈춤을 줄이는 후보였다. 4096은 평균 TTFT도 개선하지 못했다. 처리량의 약 1~2% 차이는 반복 측정 없이는 우위를 확정하기 어렵지만, ITL 꼬리 지연 차이는 후속 검증할 가치가 있다.',True)

page('4. gpu-memory-utilization 0.9와 0.8 비교')
p('기능  GPU 메모리 중 vLLM 실행기가 사용할 예산 비율이다. 가중치·실행용 메모리 등을 제외한 예산이 캐시 확보에 영향을 준다. 0.9는 KV Cache가 90% 사용 중이라는 뜻이 아니다. [1, 2]',True)
shots([(g,'seq24 · gpu-memory-utilization=0.8')],3.35)
table([(s,'0.9 기준'),(g,'0.8')])
p(f'결과  0.9→0.8에서 처리량은 {delta(g,s,"output_throughput")}, 평균 TTFT는 {delta(g,s,"mean_ttft_ms")}했다. 실행 시간은 192.86→219.81초로 늘었다. 평균 TPOT는 146.38→137.04ms로 낮아졌지만 P95 TPOT는 160.87→181.57ms로 높아졌다.',True)
p('해석과 활용  낮은 메모리 예산에서 전체 처리량과 첫 응답 대기가 악화됐다. 다만 KV 사용률·preemption 계측이 없어 캐시 부족이 원인이라고 확정할 수 없다. GPU를 단독 사용하는 이번 조건은 0.9 유지가 유리하며, 다른 프로세스에 메모리를 남겨야 할 때 0.8을 검토한다.',True)

page('5. max-model-len 4096과 8192 비교')
p('기능  요청당 입력과 출력 토큰 합계의 허용 상한이다. 값을 늘리면 더 긴 문맥을 받을 수 있지만 실제 요청이 길어지는 것은 아니다. 이번 실험은 입력 약 3,084토큰과 출력 256토큰을 그대로 유지했다. [2]',True)
shots([(l,'seq24 · max-model-len=8192')],3.35)
table([(s,'4096 기준'),(l,'8192')])
p(f'결과  출력 처리량은 132.74→132.78 token/s로 {change(l,s,"output_throughput"):.2f}% 증가했고, 평균 TTFT는 20.797→20.794초였다. P95 ITL도 827.70→827.63ms로 사실상 같았다.',True)
p('해석과 활용  허용 상한만 높여서는 이번 짧은 문맥의 처리 성능이 개선되지 않았다. KV 사용량이 두 배로 늘었다고 해석할 수도 없다. 현재 입력에는 4096이면 충분하며, 입력과 출력 합계가 4096을 넘는 장문 작업이 필요할 때 8192를 선택한다. 8192 길이 요청 자체의 성능을 검증한 실험은 아니다.',True)

page('6. enforce-eager false와 true 비교')
p('기능  true는 CUDA Graph를 사용하지 않는 Eager 실행을 강제한다. false는 CUDA Graph 등을 활용하는 일반 실행 경로를 허용한다. 최신 문서는 torch.compile 비활성화도 명시하므로 정확한 범위는 설치 버전에 따른다. [2]',True)
shots([(e,'seq24 · enforce-eager=true')],3.35)
table([(s,'false 기준'),(e,'true')])
p(f'결과  Eager 강제 시 처리량은 {delta(e,s,"output_throughput")}, 평균 TPOT는 {delta(e,s,"mean_tpot_ms")}했다. P95 TTFT는 48.30→55.67초로 {change(e,s,"p95_ttft_ms"):.1f}% 증가했고 실행 시간은 224.92초로 늘었다.',True)
p('해석과 활용  일반 실행 경로가 이번 부하에서 더 효율적이었다. 평상시에는 false를 유지하고, 그래프 실행 관련 오류를 분리하거나 호환성을 점검할 때 true를 검토한다. 이번 수치는 실행 경로 전체의 차이이며 CUDA Graph 단독 효과나 실제 메모리 절감량을 분리 측정한 것은 아니다.',True)

page('7. Prefix Caching OFF와 ON 비교')
p('기능과 조건  동일한 접두사의 계산 결과를 재사용해 입력 처리를 줄인다. seq24에서 prefix_repetition을 사용했으며, 공통 접두사 2048토큰 1종류와 요청별 접미사 1024토큰, 출력 256토큰으로 구성했다. OFF와 ON 모두 총 입력 308,504토큰으로 동일하다. [3]',True)
shots([(po,'Prefix OFF'),(pn,'Prefix ON')])
table([(po,'OFF'),(pn,'ON')])
p(f'결과  ON에서 처리량은 {delta(pn,po,"output_throughput")}, 평균 TTFT는 {delta(pn,po,"mean_ttft_ms")}했다. 실행 시간은 188.30→128.88초, P95 TTFT는 47.59→31.99초, P95 ITL은 830.81→655.21ms로 줄었다.',True)
p('해석과 활용  공통 시스템 프롬프트, 같은 문서에 대한 반복 질의, 대화 이력 재사용에 적합하다. 평균 TPOT도 145.57→98.02ms로 개선됐지만 캐시가 생성 계산을 직접 생략한 것은 아니다. 입력 처리와의 자원 경쟁 완화 가능성이 있다. 적중률은 미계측이며, 사전 단일 요청이 캐시를 예열했을 수 있어 완전한 콜드 캐시 실험으로 보지 않는다.',True)

page('8. 종합 판단과 적용 방향')
doc.add_picture(str(OUT/'comparison.png'),width=Inches(6.95))
p('그림  random 부하 8개 조건 비교. Prefix 결과는 데이터셋이 달라 이 그래프에서 제외했다.',False)
p('수치의 정상 여부  모든 요청이 완료되어 이번 단기 시험에서 요청 실패는 없었다. 그러나 random 조건의 평균 TTFT 11.51~32.06초와 평균 E2EL 57.06~66.19초는 즉각적인 채팅 응답을 기대하는 서비스에는 느린 수준이다. 긴 입력과 동시성 32의 부하 시험 결과이므로 이 값만으로 모델이나 GPU의 고장을 판정할 수는 없다.',True)
p('다른 모델과의 비교  모델 크기·구조, GPU, 입력 길이, 동시성, 엔진, 양자화가 다르면 공통 정상 범위를 적용할 수 없다. 공개 Mozilla AI Qwen3.8-27B 자료는 입력 4096에서 생성 속도 약 37.1~37.7 token/s를 제시하지만 llama.cpp·llamafile·Ollama의 생성 구간 지표이다. 이번 전체 요청 합산 출력 처리량과 직접 나누어 성능 우열을 계산할 수 없다. 같은 조건의 외부 기준이 없으므로 하드웨어 기대 성능 충족 여부는 미확정이다. [4]',True)
p('적용 후보  반복 접두사가 많으면 Prefix ON을 우선 검토한다. random 부하의 스트리밍 응답성에는 seq24·토큰 예산 1024가 후보이며, 메모리 비율 0.9·문맥 상한 4096·Eager false를 유지한다. 첫 토큰 대기 최소화가 우선이면 seq32를 선택할 수 있으나 TPOT 악화를 감수해야 한다. Prefix ON과 토큰 예산 1024의 조합은 아직 측정하지 않았다.',True)
p('과제 해석 범위  PagedAttention은 KV를 블록 단위로 관리하고, Continuous Batching은 완료된 요청의 자리에 새 요청을 넣는다. 이번 실험은 그 위의 설정 효과를 비교했으며 기능 자체의 ON/OFF 효과는 분리하지 않았다. KV 사용률·블록 수·preemption은 미계측이다. TP는 GPU 1장으로 1을 유지했고 FP8만 측정했으므로 AWQ/GPTQ 대비 성능·정확도 우위는 주장하지 않는다. Qwen3.8은 Gated DeltaNet과 Attention을 섞은 구조여서 일반 full-attention 모델의 KV 계산식을 그대로 적용할 수 없다. [5]',True)
p('해석 기준  지연은 낮을수록, 처리량은 높을수록 좋다. 변화율은 (변경값−비교 기준)/비교 기준×100이다. 각 조건 1회 결과로 장시간 안정성이나 통계적 유의성을 확정하지 않는다. 응답 품질은 이번 합성 부하 시험의 평가 대상이 아니다.',True)
refs=[('[1] vLLM 최적화','https://docs.vllm.ai/en/stable/configuration/optimization/'),('[2] vLLM 서버 옵션','https://docs.vllm.ai/en/stable/cli/serve/'),('[3] Prefix Caching','https://docs.vllm.ai/en/stable/features/automatic_prefix_caching/'),('[4] Mozilla AI 공개 결과','https://github.com/mozilla-ai/exp-llama-benchy/blob/main/data/20260912-l40s/report/qwen3.8-27b-all-servers.md'),('[5] 모델 구조','https://huggingface.co/Qwen/Qwen3.8-27B-FP8')]
pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(0)
for idx,(txt,url) in enumerate(refs):
    if idx:pp.add_run('  ·  ')
    rel=pp.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'16');pr.append(size);run.append(pr)
    tx=OxmlElement('w:t');tx.text=txt;run.append(tx);link.append(run);pp._p.append(link)

in_summary=False
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('8. 종합'):
        in_summary=True
    elif in_summary:
        paragraph.paragraph_format.line_spacing=1.05
        paragraph.paragraph_format.space_after=Pt(5)
doc.save(OUT/'vLLM_벤치마킹_결과_보고서.docx')
fields=['date','completed','failed','duration','total_input_tokens','total_output_tokens','output_throughput','mean_ttft_ms','p95_ttft_ms','mean_tpot_ms','p95_tpot_ms','p95_itl_ms','mean_e2el_ms']
(OUT/'verified_results.json').write_text(json.dumps({i:{k:data[i][k] for k in fields} for i in ids},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT/'vLLM_벤치마킹_결과_보고서.docx')

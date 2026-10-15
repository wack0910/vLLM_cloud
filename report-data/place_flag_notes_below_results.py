from pathlib import Path
import json,hashlib
from lxml import etree
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'보고서/vLLM_벤치마킹_결과_보고서_표중심_수정본.docx'
if not SRC.exists():
    SRC=ROOT/'보고서/vLLM_벤치마킹_보고서_경어체_플래그설명추가.docx'
OUT=ROOT/'보고서/vLLM_벤치마킹_보고서_결과별_플래그설명.docx'
QA=ROOT/'보고서/vllm-inline-qa';QA.mkdir(exist_ok=True)
source_hash=hashlib.sha256(SRC.read_bytes()).hexdigest()
doc=Document(SRC)
appendix=next((p._p for p in doc.paragraphs if p.text.startswith('9. 플래그 사용처')),None)
if appendix is not None:
    node=appendix
    while node is not None and node.tag!=qn('w:sectPr'):
        following=node.getnext();node.getparent().remove(node);node=following
original_paragraphs=list(doc.paragraphs)
original_xml=[etree.tostring(p._p) for p in original_paragraphs]
table_xml=[etree.tostring(t._tbl) for t in doc.tables]
changed=set();removed=set()

def replace(p,text):
    for el in list(p._p):
        if el.tag!=qn('w:pPr'):p._p.remove(el)
    if '  ' in text:
        lead,rest=text.split('  ',1);p.add_run(lead+'  ').bold=True;p.add_run(rest)
    else:p.add_run(text)

titles={
    '2. 최대 처리 요청 수 비교':'2. max-num-seqs 비교',
    '3. 최대 처리 토큰 비교':'3. max-num-batched-tokens 비교',
    '4. GPU-Mem-Utilization 값 비교':'4. gpu-memory-utilization 비교',
    '5. max-model-len 비교':'5. max-model-len 비교',
    '6. 일반 실행과 Eager 실행을 비교하면':'6. enforce-eager 비교',
    '7. 반복되는 입력을 캐시로 재사용하면':'7. enable-prefix-caching 비교',
}
replacements={
    '조정 옵션  Prefix Caching OFF / ON. -한 접두사의 입력 계산을 재사용한다. 공통 2048토큰 1종류 + 접미사 1024토큰, 출력 256토큰의 prefix_repetition 부하를 사용했다.':
    '조정 옵션  enable-prefix-caching OFF / ON을 비교했습니다. 공통 접두사 2048토큰 1종류와 접미사 1024토큰, 출력 256토큰의 prefix_repetition 부하를 사용했습니다.',
    '주요 변화  Eager 강제 시 처리량 14.25% 감소, 평균 TPOT 16.56% 증가, P95 TTFT 15.25% 증가.':
    '주요 변화  Eager 실행을 강제했을 때 처리량은 14.25% 감소했으며, 평균 TPOT는 16.56%, P95 TTFT는 15.25% 증가했습니다.',
    '주요 변화  처리량 46.11% 증가, 평균 TTFT 34.90% 감소. 같은 문서나 시스템 프롬프트를 반복하는 작업에 적합하다.':
    '주요 변화  Prefix Caching을 켰을 때 처리량은 46.11% 증가했고, 평균 TTFT는 34.90% 감소했습니다.',
    '주요 변화  처리량은 46.11% 증가했으며, 평균 TTFT는 34.90% 감소했습니다. 같은 문서나 시스템 프롬프트를 반복해서 사용하는 작업에 적합합니다.':
    '주요 변화  ON에서 처리량은 46.11% 증가했고, 평균 TTFT는 34.90% 감소했습니다.',
    '해석 범위  캐시가 생성 계산을 직접 생략하는 것은 아니다. TPOT 개선은 입력 처리와의 자원 경쟁 완화 가능성이 있다. 적중률은 미계측이며 사전 요청으로 캐시가 예열됐을 수 있다.':
    '해석 범위  입력 계산을 재사용하며 생성 계산을 직접 생략하지는 않습니다. 캐시 적중률은 측정하지 않았고, 사전 요청으로 캐시가 예열됐을 수 있습니다.',
    '해석 범위  캐시가 생성 계산을 직접 생략하지는 않습니다. TPOT는 입력 처리와의 자원 경쟁이 줄어 개선됐을 가능성이 있습니다. 적중률은 측정하지 않았으며, 사전 요청으로 캐시가 예열됐을 수 있습니다.':
    '해석 범위  입력 계산을 재사용하며 생성 계산을 직접 생략하지는 않습니다. 적중률은 측정하지 않았으며, 사전 요청으로 예열됐을 수 있습니다.',
    '처리량은 높을수록, 시간은 낮을수록 좋음 · 변화율은 반올림 전 원본값으로 계산':
    '처리량은 높을수록, 시간은 낮을수록 좋습니다. 변화율은 반올림 전 원본값으로 계산했습니다.',
    '적용 조건  Prefix ON과 토큰 예산 1024 조합은 미측정이다. 조건별 1회 결과이며 모델 품질·장시간 안정성의 검증은 아니다. Qwen3.8은 혼합 구조여서 일반 full-attention 모델의 KV 계산식을 그대로 적용하지 않는다. [5]':
    '적용 조건  Prefix ON과 토큰 예산 1024 조합은 측정하지 않았습니다. 조건별 1회 결과이므로 모델 품질과 장시간 안정성까지 검증한 것은 아닙니다. Qwen3.8은 혼합 구조이므로 일반 full-attention 모델의 KV 계산식을 그대로 적용하실 수 없습니다. [5]',
}
for idx,p in enumerate(original_paragraphs):
    if p.text in titles and titles[p.text]!=p.text:
        replace(p,titles[p.text]);changed.add(idx)
    elif p.text in replacements:
        replace(p,replacements[p.text]);changed.add(idx)
    elif p.text in [
        '활용  평상시에는 false를 유지한다. 그래프 관련 오류나 호환성을 점검할 때 true를 검토한다. 실행 경로 전체의 비교이며 CUDA Graph 단독 효과는 분리하지 않았다.',
        '활용  평상시에는 false를 유지하시는 것이 좋습니다. CUDA Graph 관련 오류나 호환성을 점검하실 때 true를 검토하실 수 있습니다. 이번 결과는 실행 경로 전체를 비교한 값이며, CUDA Graph 단독 효과는 분리하지 않았습니다.',
    ]:
        # The concise note immediately below the results now covers this point.
        p._p.getparent().remove(p._p);removed.add(idx)
    elif idx==33 and not p.text:
        # Recover a spacer line on the tightly packed context-length page.
        p._p.getparent().remove(p._p);removed.add(idx)
    elif p.text=='처리량은 높을수록, 시간은 낮을수록 좋습니다. 변화율은 반올림 전 원본값으로 계산했습니다.':
        p._p.getparent().remove(p._p);removed.add(idx)

notes={
    2:'한 단계에서 처리할 시퀀스 수를 조절합니다. 첫 응답 대기가 중요한 다중 사용자 환경에서는 큰 값을, 생성 중 응답성이 중요할 때는 작은 값을 검토하실 수 있습니다.',
    3:'단계별 토큰 처리 예산을 조절합니다. 긴 입력과 스트리밍 생성이 함께 들어오는 채팅에서는 작은 값이 생성 중 긴 멈춤을 줄이는 데 유리할 수 있습니다.',
    4:'vLLM에 배정할 GPU 메모리 예산을 정합니다. GPU 단독 사용 시에는 큰 예산을, 다른 프로세스와 공유하실 때는 낮은 예산을 검토하실 수 있습니다.',
    5:'문맥 상한을 정합니다. 긴 문서나 대화 이력을 처리할 때 늘립니다.',
    6:'Eager 실행을 강제할지 정합니다. 일반 서빙에는 false가 적합하고, CUDA Graph 관련 오류나 호환성을 점검하실 때 true를 검토하실 수 있습니다.',
    7:'공통 접두사의 입력 계산을 재사용합니다. 같은 시스템 프롬프트나 문서, 대화 이력을 반복해서 보내는 환경에서 ON이 유리합니다.',
}
def after_table(t,text):
    el=OxmlElement('w:p');t._tbl.addnext(el)
    p=Paragraph(el,t._parent)
    p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
    a=p.add_run('용도와 사용 환경  ');a.bold=True;b=p.add_run(text)
    for run in [a,b]:
        run.font.name='맑은 고딕';run.font.size=Pt(9.5);run.font.color.rgb=RGBColor(0,0,0)
        run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    return p
for table_idx,text in notes.items():after_table(doc.tables[table_idx],text)

# Preserve all existing table entries and put the unit clarification beside
# the affected result without changing the user's edited cell text.
unit=doc.add_paragraph()
unit_xml=unit._p;unit_xml.getparent().remove(unit_xml)
insert=doc.tables[3]._tbl.getnext()
insert.addnext(unit_xml)
unit=Paragraph(unit_xml,doc._body)
unit.paragraph_format.space_after=Pt(3);unit.paragraph_format.line_spacing=1
run=unit.add_run('단위 안내  위 표의 평균 TPOT 값 145.02·146.38·146.05는 ms 단위입니다. 표의 default는 비교 기준 2048을 뜻합니다.')
run.font.name='맑은 고딕';run.font.size=Pt(8);run.font.color.rgb=RGBColor(0,0,0)
run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')

doc.save(OUT)
check=Document(OUT)
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==source_hash
for idx,old in enumerate(table_xml):
    assert etree.tostring(check.tables[idx]._tbl)==old,('User table changed',idx)
# Every original node not expressly changed remains byte-identical in XML.
frozen=[old for i,old in enumerate(original_xml) if i not in changed|removed]
after_xml=[etree.tostring(p._p) for p in check.paragraphs]
for old in frozen:assert old in after_xml,'User paragraph changed or removed'
assert len(check.inline_shapes)==11
assert not any(p.text.startswith('9.') for p in check.paragraphs)
(QA/'preservation.json').write_text(json.dumps({
 'source_unchanged':True,'preserved_existing_tables':len(table_xml),
 'preserved_paragraphs':len(frozen),'changed_title_or_assistant_prose_indices':sorted(changed),
 'removed_redundant_assistant_prose_indices':sorted(removed),
 'notes_immediately_below_result_tables':len(notes),
},ensure_ascii=False,indent=2),encoding='utf-8')
print('Saved result-level notes. Existing tables and frozen paragraphs preserved.')
